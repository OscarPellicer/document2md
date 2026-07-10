from __future__ import annotations

import os
import re
import shutil
from io import BytesIO
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx"}


@dataclass(slots=True)
class ConversionOptions:
    values_only: bool = False
    include_hidden: bool = False
    include_notes: bool = False
    images: str = "placeholder"
    force: bool = False
    assets_dir: Path | None = None
    pdf_pages_per_batch: int = 10


@dataclass(slots=True)
class ConversionResult:
    input_path: Path
    output_path: Path
    assets: list[Path]
    warnings: list[str]


def _clean_inline(text: str) -> str:
    return re.sub(r"[ \t\r\f\v]+", " ", text).strip()


def _escape(text: object) -> str:
    if text is None:
        return ""
    return str(text).replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def _unique_path(path: Path, force: bool) -> Path:
    if force or not path.exists():
        return path
    index = 1
    while True:
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate
        index += 1


def _render_table(rows: list[list[object]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]

    def render(row: list[object]) -> str:
        return "| " + " | ".join(_escape(value) for value in row) + " |"

    return "\n".join(
        [render(padded[0]), "| " + " | ".join("---" for _ in range(width)) + " |"]
        + [render(row) for row in padded[1:]]
    )


def _plain_table(rows: list[list[str]]) -> str:
    """Render a table as compact text for contexts where Markdown tables cannot nest."""
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    if len(padded) == 1:
        return "; ".join(value for value in padded[0] if value)
    if width == 2:
        pairs = []
        for key, value in padded:
            if key and value:
                pairs.append(f"{key}: {value}")
            elif key or value:
                pairs.append(key or value)
        return "; ".join(pairs)
    return "; ".join(" / ".join(value for value in row if value) for row in padded)


def _iter_paragraph_runs(paragraph: Paragraph):
    for child in paragraph._p.iterchildren():
        if child.tag.endswith("}r"):
            yield child
        elif child.tag.endswith("}hyperlink"):
            yield from (grandchild for grandchild in child.iterchildren() if grandchild.tag.endswith("}r"))


def _run_text(run_element) -> str:
    texts = []
    for text_element in run_element.iter():
        if text_element.tag.endswith("}t"):
            texts.append(text_element.text or "")
        elif text_element.tag.endswith("}tab"):
            texts.append("\t")
        elif text_element.tag.endswith("}br"):
            texts.append("\n")
    return "".join(texts)


def _run_bool(run_element, name: str) -> bool | None:
    rpr = next((child for child in run_element.iterchildren() if child.tag.endswith("}rPr")), None)
    if rpr is None:
        return None
    prop = next((child for child in rpr.iterchildren() if child.tag.endswith(f"}}{name}")), None)
    if prop is None:
        return None
    value = prop.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
    return value not in {"0", "false", "False", "off"}


def _styled_runs(runs: Iterable) -> str:
    pieces: list[list[object]] = []
    for run in runs:
        if hasattr(run, "tag"):
            text = _run_text(run)
            bold = _run_bool(run, "b")
            italic = _run_bool(run, "i")
        else:
            text = run.text
            font = getattr(run, "font", None)
            bold = getattr(run, "bold", None)
            italic = getattr(run, "italic", None)
            if bold is None and font is not None:
                bold = font.bold
            if italic is None and font is not None:
                italic = font.italic
        if not text:
            continue
        style = (bool(bold), bool(italic)) if text.strip() else None
        pieces.append([text, style])

    for index, piece in enumerate(pieces):
        if piece[1] is not None:
            continue
        previous = next((part[1] for part in reversed(pieces[:index]) if part[1]), None)
        following = next((part[1] for part in pieces[index + 1 :] if part[1]), None)
        if previous == following:
            piece[1] = previous

    groups: list[list[object]] = []
    for text, style in pieces:
        style = style or (False, False)
        if groups and groups[-1][1] == style:
            groups[-1][0] = str(groups[-1][0]) + str(text)
        else:
            groups.append([text, style])

    output = []
    for text, style in groups:
        bold, italic = style
        escaped = str(text)
        if bold and italic:
            output.append(f"***{escaped}***")
        elif bold:
            output.append(f"**{escaped}**")
        elif italic:
            output.append(f"*{escaped}*")
        else:
            output.append(escaped)
    return _clean_inline("".join(output))


def _docx_paragraph(paragraph: Paragraph, in_table: bool = False) -> str:
    text = _styled_runs(_iter_paragraph_runs(paragraph))
    if not text:
        return ""
    style = paragraph.style.name if paragraph.style else ""
    heading = re.fullmatch(r"(?:Heading|Título|Encabezado)\s+([1-6])", style, re.I)
    if heading and not in_table:
        return f"{'#' * int(heading.group(1))} {text}"
    if style.lower().startswith(("list bullet", "lista con viñetas")):
        return f"- {text}"
    if style.lower().startswith(("list number", "lista numerada")):
        return f"1. {text}"
    return text


def _docx_blocks(document: DocxDocument):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield DocxTable(child, document)


def _docx_cell_blocks(cell):
    for child in cell._tc.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, cell)
        elif child.tag.endswith("}tbl"):
            yield DocxTable(child, cell)


def _docx_table_rows(table: DocxTable) -> list[list[str]]:
    rows: list[list[str]] = []
    seen = set()
    for row in table.rows:
        values = []
        for cell in row.cells:
            key = cell._tc
            if key in seen:
                values.append("")
            else:
                seen.add(key)
                values.append(_docx_cell(cell))
        rows.append(values)
    return rows


def _docx_cell(cell) -> str:
    blocks = []
    for block in _docx_cell_blocks(cell):
        if isinstance(block, Paragraph):
            markdown = _docx_paragraph(block, in_table=True)
        else:
            markdown = _plain_table(_docx_table_rows(block))
        if markdown:
            blocks.append(markdown)
    return "<br>".join(blocks)


def _docx_table(table: DocxTable) -> str:
    rows = _docx_table_rows(table)
    return _render_table(rows)


def _asset_target(assets_dir: Path, stem: str, suffix: str, used: set[str]) -> Path:
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-") or "image"
    name = f"{safe_stem}{suffix.lower()}"
    index = 1
    while name.lower() in used:
        name = f"{safe_stem}-{index}{suffix.lower()}"
        index += 1
    used.add(name.lower())
    return assets_dir / name


def _relative_link(target: Path, output_path: Path) -> str:
    return Path(os.path.relpath(target, output_path.parent)).as_posix()


def _extract_docx_images(document, output_path: Path, assets_dir: Path) -> tuple[list[str], list[Path]]:
    links = []
    assets = []
    used: set[str] = set()
    assets_dir.mkdir(parents=True, exist_ok=True)
    for index, shape in enumerate(document.inline_shapes, start=1):
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        relationship_id = blip.embed
        image_part = document.part.related_parts.get(relationship_id)
        if image_part is None:
            continue
        filename = Path(str(image_part.partname)).name
        suffix = Path(filename).suffix or ".bin"
        target = _asset_target(assets_dir, Path(filename).stem or f"image-{index}", suffix, used)
        target.write_bytes(image_part.blob)
        relative = _relative_link(target, output_path)
        links.append(f"![{target.stem}]({relative})")
        assets.append(target)
    return links, assets


def _convert_docx(path: Path, output_path: Path, options: ConversionOptions) -> tuple[str, list[Path], list[str]]:
    document = Document(str(path))
    blocks = []
    for block in _docx_blocks(document):
        markdown = _docx_paragraph(block) if isinstance(block, Paragraph) else _docx_table(block)
        if markdown:
            blocks.append(markdown)
    assets = []
    if options.images == "extract":
        assets_dir = options.assets_dir or output_path.with_name(f"{output_path.stem}_assets")
        image_links, assets = _extract_docx_images(document, output_path, assets_dir)
        if image_links:
            blocks.extend(["## Extracted images", *image_links])
    elif options.images == "embed":
        raise ValueError("--images embed is supported only for PDF inputs.")
    return "\n\n".join(blocks).rstrip() + "\n", assets, []


def _pptx_text_frame(text_frame) -> str:
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        text = _styled_runs(paragraph.runs)
        if not text:
            continue
        prefix = "- " if paragraph.level else ""
        paragraphs.append(prefix + text)
    return "\n".join(paragraphs)


def _pptx_table(table) -> str:
    rows = []
    seen = set()
    for row in table.rows:
        values = []
        for cell in row.cells:
            key = cell._tc
            if key in seen:
                values.append("")
            else:
                seen.add(key)
                values.append(_pptx_text_frame(cell.text_frame).replace("\n", "<br>"))
        rows.append(values)
    return _render_table(rows)


def _convert_pptx(path: Path, output_path: Path, options: ConversionOptions) -> tuple[str, list[Path], list[str]]:
    presentation = Presentation(str(path))
    output = []
    assets = []
    warnings = []
    assets_dir = options.assets_dir or output_path.with_name(f"{output_path.stem}_assets")
    used: set[str] = set()
    if options.images == "embed":
        raise ValueError("--images embed is supported only for PDF inputs.")

    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = _clean_inline(title_shape.text) if title_shape and title_shape.has_text_frame else ""
        output.extend([f"# Slide {slide_number}" + (f": {title}" if title else ""), ""])

        ordered_shapes = sorted(
            enumerate(slide.shapes),
            key=lambda item: (item[1].top, item[1].left, item[0]),
        )
        for _, shape in ordered_shapes:
            if shape is title_shape:
                continue
            if shape.has_table:
                output.extend([_pptx_table(shape.table), ""])
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if options.images != "extract":
                    continue
                image = shape.image
                target = _asset_target(assets_dir, f"slide-{slide_number}-image", f".{image.ext}", used)
                assets_dir.mkdir(parents=True, exist_ok=True)
                target.write_bytes(image.blob)
                assets.append(target)
                output.extend([f"![{target.stem}]({_relative_link(target, output_path)})", ""])
            elif shape.has_text_frame:
                text = _pptx_text_frame(shape.text_frame)
                if text:
                    output.extend([text, ""])
            elif shape.shape_type in {MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.GROUP}:
                warnings.append(f"Slide {slide_number}: omitted unsupported {shape.shape_type} shape.")

        if options.include_notes and slide.has_notes_slide:
            notes = _pptx_text_frame(slide.notes_slide.notes_text_frame)
            if notes:
                output.extend(["## Speaker notes", "", notes, ""])

    return "\n".join(output).rstrip() + "\n", assets, warnings


def _create_pdf_converter(options: ConversionOptions):
    try:
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import (
            PdfPipelineOptions,
            RapidOcrOptions,
            TesseractCliOcrOptions,
        )
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling_core.types.doc import ImageRefMode
    except ImportError as error:
        raise RuntimeError(
            "PDF conversion requires Docling. Install the repository dependencies "
            "or run with its .venv Python."
        ) from error

    ocr_options = (
        TesseractCliOcrOptions()
        if shutil.which("tesseract")
        else RapidOcrOptions(backend="onnxruntime")
    )
    pipeline_options = PdfPipelineOptions(
        generate_picture_images=options.images in {"extract", "embed"},
        ocr_options=ocr_options,
        ocr_batch_size=1,
        layout_batch_size=1,
        table_batch_size=1,
        queue_max_size=8,
    )
    converter = DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options),
        },
    )
    return converter, ImageRefMode


def _pdf_page_count(path: Path) -> int:
    try:
        import pypdfium2 as pdfium
    except ImportError as error:
        raise RuntimeError(
            "PDF conversion requires pypdfium2 to count pages safely."
        ) from error

    document = pdfium.PdfDocument(path.read_bytes())
    try:
        return len(document)
    finally:
        document.close()


def _pdf_stream(path: Path, batch_number: int = 0):
    """Use bytes plus an ASCII-only name to avoid native-backend path encoding bugs."""
    from docling_core.types.io import DocumentStream

    name = f"document-{batch_number:04d}.pdf" if batch_number else "document.pdf"
    return DocumentStream(name=name, stream=BytesIO(path.read_bytes()))


def _pdf_page_ranges(page_count: int, pages_per_batch: int) -> list[tuple[int, int]]:
    if pages_per_batch <= 0 or pages_per_batch >= page_count:
        return [(1, page_count)]
    return [
        (start, min(start + pages_per_batch - 1, page_count))
        for start in range(1, page_count + 1, pages_per_batch)
    ]


def _serialize_pdf_document(
    document,
    output_path: Path,
    options: ConversionOptions,
    image_ref_mode,
) -> tuple[str, list[Path], list[str]]:
    ImageRefMode = image_ref_mode

    if options.images == "extract":
        assets_dir = options.assets_dir or output_path.with_name(f"{output_path.stem}_assets")
        assets_dir.mkdir(parents=True, exist_ok=True)
        document.save_as_markdown(
            output_path,
            image_mode=ImageRefMode.REFERENCED,
            artifacts_dir=assets_dir,
        )
        assets = sorted(item for item in assets_dir.rglob("*") if item.is_file())
        return output_path.read_text(encoding="utf-8"), assets, []

    image_mode = ImageRefMode.EMBEDDED if options.images == "embed" else ImageRefMode.PLACEHOLDER
    markdown = document.export_to_markdown(image_mode=image_mode)
    return markdown.rstrip() + "\n", [], []


def _convert_pdf_with_converter(
    converter,
    image_ref_mode,
    path: Path,
    output_path: Path,
    options: ConversionOptions,
) -> tuple[str, list[Path], list[str]]:
    page_count = _pdf_page_count(path)
    if page_count < 1:
        raise RuntimeError(f"PDF has no pages: {path}")

    ranges = _pdf_page_ranges(page_count, options.pdf_pages_per_batch)
    markdown_parts: list[str] = []
    assets: list[Path] = []
    warnings: list[str] = []

    for batch_number, page_range in enumerate(ranges, start=1):
        conversion = converter.convert(
            _pdf_stream(path, batch_number),
            page_range=page_range,
        )
        document = conversion.document

        if options.images == "extract" and len(ranges) > 1:
            assets_root = options.assets_dir or output_path.with_name(f"{output_path.stem}_assets")
            chunk_assets = assets_root / f"pages-{page_range[0]:04d}-{page_range[1]:04d}"
            chunk_output = output_path.with_name(
                f".{output_path.stem}.pages-{page_range[0]:04d}-{page_range[1]:04d}.md"
            )
            chunk_options = ConversionOptions(
                images="extract",
                force=True,
                assets_dir=chunk_assets,
                pdf_pages_per_batch=options.pdf_pages_per_batch,
            )
            markdown, chunk_files, chunk_warnings = _serialize_pdf_document(
                document,
                chunk_output,
                chunk_options,
                image_ref_mode,
            )
            chunk_output.unlink(missing_ok=True)
        else:
            markdown, chunk_files, chunk_warnings = _serialize_pdf_document(
                document,
                output_path,
                options,
                image_ref_mode,
            )

        markdown_parts.append(markdown.strip())
        assets.extend(chunk_files)
        warnings.extend(chunk_warnings)

    return "\n\n".join(part for part in markdown_parts if part).rstrip() + "\n", assets, warnings


def _convert_pdf(path: Path, output_path: Path, options: ConversionOptions) -> tuple[str, list[Path], list[str]]:
    converter, image_ref_mode = _create_pdf_converter(options)
    return _convert_pdf_with_converter(converter, image_ref_mode, path, output_path, options)


def _xlsx_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _trim_sheet_rows(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(cell != "" for cell in rows[-1]):
        rows.pop()
    if not rows:
        return []
    last_column = max(
        (index for row in rows for index, cell in enumerate(row) if cell != ""),
        default=-1,
    )
    return [row[: last_column + 1] for row in rows]


def _convert_xlsx(path: Path, output_path: Path, options: ConversionOptions) -> tuple[str, list[Path], list[str]]:
    workbook = load_workbook(path, data_only=options.values_only, read_only=False)
    output = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible" and not options.include_hidden:
            continue
        rows = [
            [_xlsx_value(cell.value) for cell in row]
            for row in sheet.iter_rows(
                min_row=1,
                max_row=sheet.max_row,
                min_col=1,
                max_col=sheet.max_column,
            )
        ]
        rows = _trim_sheet_rows(rows)
        if not rows:
            continue
        output.extend([f"# Sheet: {sheet.title}", "", _render_table(rows), ""])
    return "\n".join(output).rstrip() + "\n", [], []


def convert_file(input_path: str | Path, output_path: str | Path | None = None, options: ConversionOptions | None = None) -> ConversionResult:
    options = options or ConversionOptions()
    source = Path(input_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Input file does not exist: {source}")
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported input type {suffix!r}; expected PDF, DOCX, PPTX, or XLSX.")
    if options.images not in {"placeholder", "extract", "embed"}:
        raise ValueError("images must be one of: placeholder, extract, embed.")
    if options.pdf_pages_per_batch < 0:
        raise ValueError("pdf_pages_per_batch must be zero or a positive integer.")

    requested = Path(output_path).expanduser() if output_path else source.with_suffix(".md")
    target = _unique_path(requested.resolve(), options.force)
    target.parent.mkdir(parents=True, exist_ok=True)

    if suffix == ".pdf":
        markdown, assets, warnings = _convert_pdf(source, target, options)
    elif suffix == ".docx":
        markdown, assets, warnings = _convert_docx(source, target, options)
    elif suffix == ".pptx":
        markdown, assets, warnings = _convert_pptx(source, target, options)
    else:
        markdown, assets, warnings = _convert_xlsx(source, target, options)

    if not target.exists() or suffix != ".pdf" or options.images != "extract":
        target.write_text(markdown, encoding="utf-8", newline="\n")
    return ConversionResult(source, target, assets, warnings)


def convert_files(
    input_paths: Iterable[str | Path],
    output_paths: Iterable[str | Path],
    options: ConversionOptions | None = None,
) -> list[ConversionResult]:
    """Convert several files, reusing one Docling converter for all PDFs."""
    options = options or ConversionOptions()
    pairs = list(zip(input_paths, output_paths, strict=True))
    results: list[ConversionResult | None] = [None] * len(pairs)
    pdf_jobs: list[tuple[int, Path, Path]] = []

    for index, (raw_input, raw_output) in enumerate(pairs):
        source = Path(raw_input).expanduser().resolve()
        if source.suffix.lower() == ".pdf":
            if not source.is_file():
                raise FileNotFoundError(f"Input file does not exist: {source}")
            target = _unique_path(Path(raw_output).expanduser().resolve(), options.force)
            target.parent.mkdir(parents=True, exist_ok=True)
            pdf_jobs.append((index, source, target))
        else:
            results[index] = convert_file(raw_input, raw_output, options)

    if pdf_jobs:
        converter, image_ref_mode = _create_pdf_converter(options)
        try:
            for index, source, target in pdf_jobs:
                markdown, assets, warnings = _convert_pdf_with_converter(
                    converter,
                    image_ref_mode,
                    source,
                    target,
                    options,
                )
                if not target.exists() or options.images != "extract":
                    target.write_text(markdown, encoding="utf-8", newline="\n")
                results[index] = ConversionResult(source, target, assets, warnings)
        except Exception as error:
            raise RuntimeError(f"Docling PDF conversion failed: {error}") from error

    return [result for result in results if result is not None]
