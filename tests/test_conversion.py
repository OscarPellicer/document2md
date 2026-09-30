from __future__ import annotations

import subprocess
import sys
import os
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from document2md import converter as converter_module
from document2md.converter import ConversionOptions, convert_file, convert_files


ROOT = Path(__file__).resolve().parents[1]
SKILL_SCRIPT = ROOT / "skills" / "document2md" / "scripts" / "convert_document.py"


def test_docx_preserves_paragraph_table_order(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Real heading", level=1)
    document.add_paragraph("Before table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    document.add_paragraph("After table.")
    source = tmp_path / "sample.docx"
    document.save(source)

    result = convert_file(source, options=ConversionOptions(force=True))
    markdown = result.output_path.read_text(encoding="utf-8")

    assert markdown.index("# Real heading") < markdown.index("Before table.")
    assert "| A | B |" in markdown
    assert markdown.index("| A | B |") < markdown.index("After table.")


def test_docx_does_not_infer_bold_heading_and_keeps_merged_table_grid(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Bold but not a heading").bold = True
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(0, 2).text = "Third"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    table.cell(1, 2).text = "C"
    source = tmp_path / "merged.docx"
    document.save(source)

    markdown = convert_file(source, options=ConversionOptions(force=True)).output_path.read_text(encoding="utf-8")
    assert "**Bold but not a heading**" in markdown
    assert "# Bold but not a heading" not in markdown
    assert "| Merged |  | Third |" in markdown


def test_docx_renders_nested_tables_inside_cells_as_compact_text(tmp_path: Path) -> None:
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.text = "Before nested table."
    nested = cell.add_table(rows=3, cols=2)
    nested.cell(0, 0).text = "Concept"
    nested.cell(0, 1).text = "Weight"
    nested.cell(1, 0).text = "Continuous assessment"
    nested.cell(1, 1).text = "30%"
    nested.cell(2, 0).text = "Final exam"
    nested.cell(2, 1).text = "70%"
    cell.add_paragraph("After nested table.")
    source = tmp_path / "nested-table.docx"
    document.save(source)

    markdown = convert_file(source, options=ConversionOptions(force=True)).output_path.read_text(encoding="utf-8")

    assert "Before nested table." in markdown
    assert "Concept: Weight; Continuous assessment: 30%; Final exam: 70%" in markdown
    assert "\\| Concept" not in markdown
    assert markdown.index("Before nested table.") < markdown.index("Concept")
    assert markdown.index("70%") < markdown.index("After nested table.")


def test_docx_includes_text_inside_hyperlinks(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Regulation: ")
    part = paragraph.part
    rel_id = part.relate_to(
        "https://example.test/regulation",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "https://example.test/regulation"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    source = tmp_path / "linked.docx"
    document.save(source)

    markdown = convert_file(source, options=ConversionOptions(force=True)).output_path.read_text(encoding="utf-8")

    assert "Regulation: https://example.test/regulation" in markdown


def test_docx_images_are_only_saved_when_requested(tmp_path: Path) -> None:
    image_path = tmp_path / "figure.png"
    Image.new("RGB", (20, 20), "red").save(image_path)
    document = Document()
    document.add_paragraph("Text")
    document.add_picture(str(image_path))
    source = tmp_path / "pictured.docx"
    document.save(source)

    placeholder_output = tmp_path / "placeholder.md"
    placeholder = convert_file(
        source,
        placeholder_output,
        ConversionOptions(images="placeholder", force=True),
    )
    assert placeholder.assets == []
    assert not (tmp_path / "placeholder_assets").exists()

    extracted_output = tmp_path / "extracted.md"
    extracted = convert_file(
        source,
        extracted_output,
        ConversionOptions(images="extract", force=True),
    )
    assert len(extracted.assets) == 1
    assert extracted.assets[0].is_file()
    assert "extracted_assets/" in extracted_output.read_text(encoding="utf-8")


def test_pptx_converts_slide_text_and_table(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Overview"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1))
    box.text = "Body text"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Count"
    table.cell(1, 1).text = "3"
    source = tmp_path / "slides.pptx"
    presentation.save(source)

    markdown = convert_file(source, options=ConversionOptions(force=True)).output_path.read_text(encoding="utf-8")
    assert "# Slide 1: Overview" in markdown
    assert "Body text" in markdown
    assert "| Metric | Value |" in markdown
    assert markdown.count("Overview") == 1


def test_xlsx_preserves_formulas_and_sheet_boundaries(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Item", "Value"])
    sheet.append(["Total", "=1+2"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Secret"
    source = tmp_path / "book.xlsx"
    workbook.save(source)

    markdown = convert_file(source, options=ConversionOptions(force=True)).output_path.read_text(encoding="utf-8")
    assert "# Sheet: Data" in markdown
    assert "=1+2" in markdown
    assert "Hidden" not in markdown


def test_skill_script_is_runnable_without_package_install(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.append(["A", "B"])
    source = tmp_path / "book.xlsx"
    workbook.save(source)

    result = subprocess.run(
        [sys.executable, str(SKILL_SCRIPT), str(source), "--force"],
        check=True,
        capture_output=True,
        text=True,
    )
    assert str(tmp_path / "book.md") in result.stdout
    assert "| A | B |" in (tmp_path / "book.md").read_text(encoding="utf-8")


def test_skill_script_prints_unicode_paths_on_legacy_console(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.append(["A", "B"])
    source = tmp_path / "Matema\u0301ticas II.xlsx"
    workbook.save(source)
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "cp1252"

    result = subprocess.run(
        [sys.executable, str(SKILL_SCRIPT), str(source), "--force"],
        check=True,
        capture_output=True,
        env=env,
    )

    assert "Matema" in result.stdout.decode("utf-8")
    assert (tmp_path / "Matema\u0301ticas II.md").exists()


def test_multiple_pdfs_share_one_docling_converter(tmp_path: Path, monkeypatch) -> None:
    sources = [tmp_path / "one.pdf", tmp_path / "two.pdf"]
    for source in sources:
        source.write_bytes(b"%PDF-test")
    targets = [tmp_path / "one.md", tmp_path / "two.md"]
    created = []
    converted_paths = []

    class FakeConverter:
        def convert(self, source, *, page_range):
            converted_paths.append((source.name, page_range))
            name = "one" if len(converted_paths) == 1 else "two"
            return SimpleNamespace(document=SimpleNamespace(name=name))

    def fake_create(options):
        created.append(options)
        return FakeConverter(), object()

    def fake_serialize(document, output_path, options, image_ref_mode):
        return f"# {document.name}\n", [], []

    monkeypatch.setattr(converter_module, "_create_pdf_converter", fake_create)
    monkeypatch.setattr(converter_module, "_serialize_pdf_document", fake_serialize)
    monkeypatch.setattr(converter_module, "_pdf_page_count", lambda path: 1)

    batch = convert_files(sources, targets, ConversionOptions(force=True))
    results = batch.results

    assert batch.failures == []
    assert len(created) == 1
    assert converted_paths == [
        ("document-0001.pdf", (1, 1)),
        ("document-0001.pdf", (1, 1)),
    ]
    assert [result.output_path for result in results] == targets
    assert targets[0].read_text(encoding="utf-8") == "# one\n"
    assert targets[1].read_text(encoding="utf-8") == "# two\n"


def test_pdf_is_processed_in_ordered_page_ranges_with_ascii_stream_names(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "memória verificación.pdf"
    source.write_bytes(b"%PDF-test")
    output = tmp_path / "result.md"
    calls = []

    class FakeDocument:
        def __init__(self, page_range):
            self.page_range = page_range

        def export_to_markdown(self, *, image_mode):
            return f"pages {self.page_range[0]}-{self.page_range[1]}"

    class FakeConverter:
        def convert(self, stream, *, page_range):
            calls.append((stream.name, page_range, stream.stream.read()))
            return SimpleNamespace(document=FakeDocument(page_range))

    monkeypatch.setattr(
        converter_module,
        "_create_pdf_converter",
        lambda options: (
            FakeConverter(),
            SimpleNamespace(PLACEHOLDER="placeholder", EMBEDDED="embedded"),
        ),
    )
    monkeypatch.setattr(converter_module, "_pdf_page_count", lambda path: 23)

    convert_file(
        source,
        output,
        ConversionOptions(force=True, pdf_pages_per_batch=10),
    )

    assert calls == [
        ("document-0001.pdf", (1, 10), b"%PDF-test"),
        ("document-0002.pdf", (11, 20), b"%PDF-test"),
        ("document-0003.pdf", (21, 23), b"%PDF-test"),
    ]
    assert output.read_text(encoding="utf-8") == (
        "pages 1-10\n\npages 11-20\n\npages 21-23\n"
    )


def test_pdf_page_ranges_can_disable_batching() -> None:
    assert converter_module._pdf_page_ranges(101, 0) == [(1, 101)]
    assert converter_module._pdf_page_ranges(5, 10) == [(1, 5)]


def test_pdf_extract_mode_uses_accompanying_artifacts_folder(tmp_path: Path) -> None:
    class FakeDocument:
        """Mimics docling_core: a relative artifacts_dir resolves against the Markdown."""

        def save_as_markdown(self, filename, *, image_mode, artifacts_dir):
            resolved = artifacts_dir if artifacts_dir.is_absolute() else filename.parent / artifacts_dir
            resolved.mkdir(parents=True, exist_ok=True)
            (resolved / "picture.png").write_bytes(b"png")
            filename.write_text("![picture](sample_assets/picture.png)\n", encoding="utf-8")

    output = tmp_path / "sample.md"
    markdown, assets, warnings = converter_module._serialize_pdf_document(
        FakeDocument(),
        output,
        ConversionOptions(images="extract"),
        SimpleNamespace(REFERENCED="referenced"),
    )

    assert warnings == []
    assert markdown.startswith("![picture]")
    assert assets == [tmp_path / "sample_assets" / "picture.png"]
